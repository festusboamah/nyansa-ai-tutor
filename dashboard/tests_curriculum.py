import copy
import json
from datetime import date
from io import BytesIO
from unittest.mock import patch

from django.test import SimpleTestCase, TestCase
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from docx import Document

from accounts.models import User
from courses.models import Subject
from schools.models import School, SchoolMembership
from .curriculum import basic_level, codes, curriculum_pages, valid_codes, valid_alignment, source_wording
from .forms import SchemeOfLearningForm
from .models import LessonNote, SchemeOfLearning
from .scheme_ai import generate_scheme_of_learning
from .lesson_ai import generate_lesson_note
from .lesson_docx import build_lesson_note_docx
from .scheme_docx import build_scheme_of_learning_docx


STANDARD = "B7.3.1.1: Identify and Explain the Importance of the Family Systems"
INDICATOR = "B7.3.1.1.1: Explain the concept and types of family systems in Ghana"


def term_week(number=1):
    return dict(week=number, strand="The family and the community", sub_strand="Family systems",
                content_standard=STANDARD, indicators=INDICATOR, resources="Family tree pictures")


class CurriculumGenerationTests(SimpleTestCase):
    def test_aliases_and_real_pack_references(self):
        self.assertEqual(basic_level("JHS 2"), 8)
        self.assertEqual(basic_level("Basic 7"), 7)
        self.assertEqual(basic_level("B.S.7"), 7)
        self.assertIsNone(basic_level("JHS 9"))
        pages = curriculum_pages("R.M.E", "B7", INDICATOR)
        self.assertIn(41, [p["page"] for p in pages])
        self.assertTrue(valid_codes(INDICATOR, pages, "B7"))
        self.assertFalse(valid_codes("B8.3.1.1.1", pages, "B7"))
        self.assertEqual(codes("B7 3.1.1"), {"B7.3.1.1"})
        self.assertFalse(valid_alignment(STANDARD, "B7.4.1.1.1"))
        self.assertTrue(source_wording(STANDARD, pages))
        self.assertFalse(source_wording("B7.3.1.1: Invented description", pages))

    def test_primary_and_unknown_subject_do_not_borrow_jhs_codes(self):
        self.assertEqual(curriculum_pages("RME", "Basic 3"), [])
        self.assertEqual(curriculum_pages("Unknown", "B7"), [])

    @patch("dashboard.scheme_ai.complete_json")
    def test_termly_repeated_strands_and_grounded_codes(self, complete):
        complete.return_value = {"weeks": [term_week(1), term_week(2)]}
        result = generate_scheme_of_learning("B7", "RME", "Term 2", 2)
        self.assertEqual(result["weeks"][0]["strand"], result["weeks"][1]["strand"])
        self.assertTrue(result["curriculum_sources"])
        self.assertIn("untrusted reference DATA", complete.call_args.args[0])

    @patch("dashboard.scheme_ai.complete_json")
    def test_invalid_or_unreferenced_scheme_is_rejected(self, complete):
        bad = [None, [], {"weeks": []}, {"weeks": ["bad"]}, {"weeks": [{"week": 1, "topic": "Family"}]}]
        for field, value in [("week", True), ("week", 2), ("indicators", "B7.99.99.99.99"), ("resources", "")]:
            row = term_week()
            row[field] = value
            bad.append({"weeks": [row]})
        for result in bad:
            with self.subTest(result=result):
                complete.return_value = result
                self.assertIsNone(generate_scheme_of_learning("B7", "RME", "Term 2", 1))

    @patch("dashboard.scheme_ai.complete_json")
    def test_yearly_has_three_terms_and_allows_multiweek_topics(self, complete):
        complete.return_value = {"weeks": [dict(week=n, term_1="Worship", term_2="Family systems", term_3="Manners") for n in (1, 2)]}
        result = generate_scheme_of_learning("B7", "RME", "All three terms", 2, plan_type="YEARLY", academic_year="2026/2027")
        self.assertEqual(result["plan_type"], "YEARLY")
        self.assertEqual(result["weeks"][1]["term_2"], "Family systems")
        complete.return_value["weeks"][0].pop("term_3")
        self.assertIsNone(generate_scheme_of_learning("B7", "RME", "All three terms", 2, plan_type="YEARLY"))

    @patch("dashboard.lesson_ai.complete_json")
    def test_lesson_uses_source_and_rejects_invented_codes_and_wrong_days(self, complete):
        payload = dict(content_standard=STANDARD, learning_indicator=INDICATOR, performance_indicators="Explain family systems",
                       core_competencies="Communication", resources="Pictures",
                       days=[dict(day="Monday", starter="Discuss families", main="Draw a family tree", reflection="Explain relationships")])
        kwargs = dict(class_level="B7", subject_name="RME", week_number=2, week_ending=date(2026, 1, 16), strand_topic="Family systems",
                      content_standard="", learning_indicator="", performance_indicator="", reference="", resources="", teaching_days=["Monday"])
        complete.return_value = copy.deepcopy(payload)
        result = generate_lesson_note(**kwargs)
        self.assertEqual(result["week"], "Week 2")
        self.assertIn(41, [p["page"] for p in result["curriculum_sources"]])
        complete.return_value["learning_indicator"] = "B7.99.99.99.99: Invented"
        self.assertIsNone(generate_lesson_note(**kwargs))
        complete.return_value = copy.deepcopy(payload)
        complete.return_value["days"][0]["day"] = "Tuesday"
        self.assertIsNone(generate_lesson_note(**kwargs))


class PlanningOutputTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.school = School.objects.create(name="Planning School", slug="planning-school")
        cls.teacher = User.objects.create_user(username="planner", first_name="Ama", last_name="Mensah", role="TEACHER")
        SchoolMembership.objects.create(school=cls.school, user=cls.teacher, role="TEACHER")
        cls.subject = Subject.objects.create(school=cls.school, name="RME")

    def test_yearly_form_requires_year_and_bounds_weeks(self):
        data = dict(subject=self.subject.pk, class_level="B7", plan_type="YEARLY", num_weeks=12)
        self.assertFalse(SchemeOfLearningForm(data, school=self.school).is_valid())
        data["academic_year"] = "2026/2027"
        form = SchemeOfLearningForm(data, school=self.school)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["term"], "All three terms")
        data["num_weeks"] = 1000
        self.assertFalse(SchemeOfLearningForm(data, school=self.school).is_valid())

    def test_cross_school_subject_rejected(self):
        other = School.objects.create(name="Other", slug="other-planning")
        form = SchemeOfLearningForm(dict(subject=self.subject.pk, class_level="B7", term="Term 2", num_weeks=12), school=other)
        self.assertFalse(form.is_valid())
        self.assertIn("subject", form.errors)

    def test_word_exports_for_both_plan_types_and_legacy(self):
        scheme = SchemeOfLearning(teacher=self.teacher, subject=self.subject, class_level="B7", term="Term 2")
        doc = Document(build_scheme_of_learning_docx(scheme, {"weeks": [term_week()]}))
        self.assertEqual([c.text for c in doc.tables[0].rows[0].cells], ["Week", "Strand", "Sub-Strand", "Content Standard", "Indicators", "Resources"])
        self.assertIn(INDICATOR, [c.text for c in doc.tables[0].rows[1].cells])
        scheme.plan_type = "YEARLY"
        doc = Document(build_scheme_of_learning_docx(scheme, {"weeks": [dict(week=1, term_1="Worship", term_2="Family", term_3="Manners")]}))
        self.assertEqual(len(doc.tables[0].columns), 4)
        scheme.plan_type = "TERMLY"
        doc = Document(build_scheme_of_learning_docx(scheme, {"weeks": [dict(week=1, topic="Legacy topic")]}))
        self.assertEqual(doc.tables[0].rows[1].cells[1].text, "Legacy topic")

    def test_lesson_signoff_does_not_claim_draft_or_return_was_vetted(self):
        note = LessonNote(teacher=self.teacher, teacher_name="Ama Mensah", subject=self.subject, class_level="B7",
                          week_number=3, strand_topic="Family", week_ending=date(2026, 1, 16), reviewed_at=timezone.now())
        for state in ["DRAFT", "SENT_BACK", "PENDING_REVIEW"]:
            note.status = state
            self.assertIsNone(note.date_vetted)
        note.status = "APPROVED"
        self.assertEqual(note.date_vetted, note.reviewed_at)
        doc = Document(build_lesson_note_docx(note, {"days": []}))
        text = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        for expected in ["Teacher Name", "Ama Mensah", "Week 3", "Date Vetted", "Headteacher Signature"]:
            self.assertIn(expected, text)
        html = render_to_string("dashboard/lesson_note_pdf.html", dict(note=note, lesson_data={"days": []}))
        self.assertIn("Ama Mensah", html)
        self.assertIn("Week 3", html)
        self.assertIn("Date Vetted", html)

    def test_teacher_name_is_versioned_and_locked_after_approval(self):
        from django.core.exceptions import ValidationError
        from .lesson_workflow import record_initial_lesson_version
        note = LessonNote.objects.create(teacher=self.teacher, teacher_name="Ama Mensah", subject=self.subject,
            class_level="B7", week_number=2, strand_topic="Family", week_ending=date(2026, 1, 16), generated_content='{"days": []}')
        actor = SchoolMembership.objects.get(user=self.teacher, school=self.school)
        version = record_initial_lesson_version(note=note, actor=actor)
        self.assertEqual(version.snapshot["teacher_name"], "Ama Mensah")
        self.assertEqual(version.snapshot["week_number"], 2)
        note.status = "APPROVED"
        note.save()
        note.teacher_name = "Different name"
        with self.assertRaises(ValidationError):
            note.save()

    def test_preview_escapes_generated_resources(self):
        self.client.force_login(self.teacher)
        row = term_week()
        row["resources"] = '<script>alert("x")</script>'
        scheme = SchemeOfLearning.objects.create(teacher=self.teacher, subject=self.subject,
            class_level="B7", term="Term 2", generated_content=json.dumps({"weeks": [row]}))
        response = self.client.get(reverse("scheme_of_learning_detail", args=[scheme.pk]), secure=True)
        self.assertContains(response, "&lt;script&gt;")
        self.assertNotContains(response, '<script>alert("x")</script>')

    @patch("dashboard.views.generate_scheme_of_learning")
    def test_yearly_route_persists_scope_and_protects_access(self, generate):
        generate.return_value = {"weeks": [dict(week=1, term_1="Worship", term_2="Family", term_3="Manners")]}
        self.client.force_login(self.teacher)
        response = self.client.post(reverse("create_scheme_of_learning"), dict(subject=self.subject.pk, class_level="B7",
            plan_type="YEARLY", academic_year="2026/2027", num_weeks=1), secure=True)
        self.assertEqual(response.status_code, 302)
        scheme = SchemeOfLearning.objects.get()
        self.assertEqual(scheme.plan_type, "YEARLY")
        self.assertEqual(generate.call_args.kwargs["academic_year"], "2026/2027")
        response = self.client.get(reverse("scheme_of_learning_detail", args=[scheme.pk]), secure=True)
        self.assertContains(response, "Third Term")
        other = User.objects.create_user(username="other-planner", role="TEACHER")
        SchoolMembership.objects.create(school=self.school, user=other, role="TEACHER")
        self.client.force_login(other)
        for route in ["scheme_of_learning_detail", "download_scheme_of_learning_docx"]:
            self.assertEqual(self.client.get(reverse(route, args=[scheme.pk]), secure=True).status_code, 404)
